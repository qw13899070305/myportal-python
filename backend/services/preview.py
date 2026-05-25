import os, io, base64
from pathlib import Path
from fastapi import HTTPException
from backend.core.config import settings

def get_file_path(filename: str) -> Path:
    # ✅ 防御路径穿越
    safe_name = os.path.basename(filename)
    fp = (settings.UPLOAD_DIR / safe_name).resolve()
    if not str(fp).startswith(str(settings.UPLOAD_DIR.resolve())):
        raise HTTPException(400, "非法文件路径")
    if not fp.exists():
        raise HTTPException(404, "文件不存在")
    return fp

def preview_word(file_path: Path) -> str:
    try:
        import mammoth
        with open(file_path, "rb") as f:
            result = mammoth.convert_to_html(f)
        return f"<div class='preview-container'>{result.value}</div>"
    except ImportError:
        from docx import Document
        doc = Document(str(file_path))
        parts = ["<div class='preview-container'>"]
        for para in doc.paragraphs:
            text = para.text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            parts.append(f"<p>{text}</p>")
        parts.append("</div>")
        return "".join(parts)

def preview_excel(file_path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(file_path), read_only=True, data_only=True)
    html = ["<div class='preview-container'>"]
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        html.append(f"<h3>{sheet_name}</h3><table border='1'>")
        for row in ws.iter_rows(values_only=True):
            html.append("<tr>")
            for cell in row:
                val = str(cell) if cell is not None else ""
                html.append(f"<td>{val}</td>")
            html.append("</tr>")
        html.append("</table>")
    html.append("</div>")
    return "".join(html)

def preview_ppt(file_path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(file_path))
    html = ["<div class='preview-container'>"]
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
        content = "<br>".join(texts)
        html.append(f"<h2>第 {i} 页</h2><p>{content}</p>")
    html.append("</div>")
    return "".join(html)

def preview_pdf_page(file_path: Path, page_num: int = 1) -> bytes:
    import fitz
    with fitz.open(str(file_path)) as doc:
        if page_num < 1 or page_num > doc.page_count:
            raise HTTPException(404, "页码超出范围")
        page = doc.load_page(page_num - 1)
        pix = page.get_pixmap(dpi=150)
        return pix.tobytes("jpg")

def get_pdf_page_count(file_path: Path) -> int:
    import fitz
    with fitz.open(str(file_path)) as doc:
        return doc.page_count
