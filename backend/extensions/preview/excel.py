"""Excel（.xlsx）预览扩展：逐工作表转成 HTML 表格。"""

from __future__ import annotations

from pathlib import Path

from backend.extensions.preview.base import Preview, PreviewHandler
from backend.extensions.registry import get_registry


class ExcelPreview(PreviewHandler):
    name = "excel"
    extensions = frozenset({".xlsx"})
    priority = 10

    #: 单个工作表最多渲染的行数，避免超大表格拖垮浏览器
    max_rows = 500

    def render(self, path: Path) -> Preview:
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=True)
        parts = ["<div class='preview-container'>"]
        try:
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                parts.append(f"<h3>{self.escape(sheet_name)}</h3><table>")
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= self.max_rows:
                        parts.append(
                            "<tr><td colspan='99'>…（表格过大，仅显示前 "
                            f"{self.max_rows} 行）</td></tr>"
                        )
                        break
                    parts.append("<tr>")
                    for cell in row:
                        value = "" if cell is None else self.escape(cell)
                        parts.append(f"<td>{value}</td>")
                    parts.append("</tr>")
                parts.append("</table>")
        finally:
            workbook.close()
        parts.append("</div>")
        return Preview.html_page(self.wrap_html("".join(parts)))


get_registry("preview").register(ExcelPreview())
