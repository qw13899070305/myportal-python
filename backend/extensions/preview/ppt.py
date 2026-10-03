"""PowerPoint（.pptx）预览扩展：逐页提取文本框。"""

from __future__ import annotations

from pathlib import Path

from backend.extensions.preview.base import Preview, PreviewHandler
from backend.extensions.registry import get_registry


class PptPreview(PreviewHandler):
    name = "ppt"
    extensions = frozenset({".pptx"})
    priority = 10

    def render(self, path: Path) -> Preview:
        from pptx import Presentation

        presentation = Presentation(str(path))
        parts = ["<div class='preview-container'>"]
        for index, slide in enumerate(presentation.slides, 1):
            texts: list[str] = []
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        texts.append(self.escape(text))
            body = "<br>".join(texts) if texts else "<em>（本页无文本内容）</em>"
            parts.append(f"<h3>第 {index} 页</h3><p>{body}</p>")
        parts.append("</div>")
        return Preview.html_page(self.wrap_html("".join(parts)))


get_registry("preview").register(PptPreview())
